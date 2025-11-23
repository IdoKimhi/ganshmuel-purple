import os
import io
from datetime import datetime
from typing import Optional, Dict, Any, List

from flask import Flask, request, jsonify, send_file
import pymysql
from pymysql.err import IntegrityError
import requests
from openpyxl import load_workbook, Workbook

# -------------------------------------------------------------------
# Configuration
# -------------------------------------------------------------------

DB_CONFIG = {
    "host": os.environ.get("DB_HOST", "localhost"),
    "user": os.environ.get("DB_USER", "root"),
    "password": os.environ.get("DB_PASSWORD", "12345678"),
    "database": os.environ.get("DB_NAME", "billdb"),
    "port": int(os.environ.get("DB_PORT", 3306)),
}

# Base URL of Weight service, e.g. "http://weight:5000"
WEIGHT_SERVICE_URL = os.environ.get("WEIGHT_SERVICE_URL", "http://weight:5000")

# Directory inside container where rates files live (mounted volume)
RATES_DIR = os.environ.get("RATES_DIR", "/in")


def create_app() -> Flask:
    app = Flask(__name__)

    # ---------------------------------------------------------------
    # Helpers
    # ---------------------------------------------------------------

    def get_db_connection():
        return pymysql.connect(
            host=DB_CONFIG["host"],
            user=DB_CONFIG["user"],
            password=DB_CONFIG["password"],
            database=DB_CONFIG["database"],
            port=DB_CONFIG["port"],
            cursorclass=pymysql.cursors.DictCursor,
            autocommit=True,
        )

    def parse_datetime_param(name: str) -> str:
        """
        Parse t1/t2 from query: yyyymmddhhmmss.
        Defaults:
        - from: 1st of month at 000000
        - to:   now
        """
        value = request.args.get(name)
        if value:
            return value

        now = datetime.now()
        if name == "from":
            dt = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        else:
            dt = now
        return dt.strftime("%Y%m%d%H%M%S")

    def call_weight_service(path: str, params: Optional[Dict[str, Any]] = None) -> Any:
        """
        Call Weight microservice (GET only, as per spec).
        """
        url = WEIGHT_SERVICE_URL.rstrip("/") + path
        resp = requests.get(url, params=params, timeout=5)
        resp.raise_for_status()
        return resp.json()

    # ---------------------------------------------------------------
    # GET /health
    # ---------------------------------------------------------------
    @app.route("/health", methods=["GET"])
    def health():
        return "OK (placeholder)", 200

    # ---------------------------------------------------------------
    # Providers
    # ---------------------------------------------------------------
    @app.route("/provider", methods=["POST"])
    def create_provider():
        return jsonify({"id": "placeholder"}), 201

    @app.route("/provider/<int:provider_id>", methods=["PUT"])
    def update_provider(provider_id: int):
        return jsonify({"id": provider_id, "name": "placeholder"}), 200

    # ---------------------------------------------------------------
    # Rates
    # ---------------------------------------------------------------

    def replace_all_rates(rates: List[Dict[str, Any]]) -> None:
        """
        Delete all rows from Rates and insert the provided ones.
        Each item: { "product_id": str, "rate": int, "scope": str }
        """
        conn = get_db_connection()
        try:
            with conn.cursor() as cur:
                cur.execute("DELETE FROM Rates")
                sql = "INSERT INTO Rates (product_id, rate, scope) VALUES (%s, %s, %s)"
                for r in rates:
                    cur.execute(sql, (r["product_id"], r["rate"], r["scope"]))
        finally:
            conn.close()

    @app.route("/rates", methods=["POST"])
    def upload_rates():
        """
        POST /rates
        - file=<filename> (query string)
        Reads Excel from "/in/<filename>" inside container.
        Rate Excel must have columns: Product, Rate, Scope
        - Product : product id (e.g. "Navel")
        - Rate    : integer (agorot)
        - Scope   : "ALL" or provider id (as string)
        New rates overwrite old ones.
        """
        filename = request.args.get("file")
        if not filename:
            return jsonify({"error": "file parameter is required, e.g. ?file=rates.xlsx"}), 400

        rates_path = os.path.join(RATES_DIR, filename)
        if not os.path.isfile(rates_path):
            return jsonify({"error": "file not found in /in", "path": rates_path}), 404

        try:
            wb = load_workbook(rates_path, data_only=True)
            ws = wb.active

            header_row = next(ws.iter_rows(min_row=1, max_row=1))
            headers = [
                (str(cell.value).strip() if cell.value is not None else "")
                for cell in header_row
            ]
            header_map = {h.lower(): i for i, h in enumerate(headers)}

            required = ["product", "rate", "scope"]
            for col in required:
                if col not in header_map:
                    return jsonify({"error": f"missing '{col}' column in header"}), 400

            rates: List[Dict[str, Any]] = []
            for row in ws.iter_rows(min_row=2):
                product_cell = row[header_map["product"]].value
                rate_cell = row[header_map["rate"]].value
                scope_cell = row[header_map["scope"]].value

                # skip empty rows
                if product_cell is None and rate_cell is None and scope_cell is None:
                    continue

                product_id = str(product_cell).strip()
                try:
                    rate_int = int(rate_cell)
                except (TypeError, ValueError):
                    return jsonify({"error": f"invalid rate value: {rate_cell!r}"}), 400

                scope = "ALL" if scope_cell is None else str(scope_cell).strip()

                rates.append(
                    {
                        "product_id": product_id,
                        "rate": rate_int,
                        "scope": scope,
                    }
                )

            replace_all_rates(rates)
        except Exception as e:
            return jsonify({"error": "failed to process rates file", "details": str(e)}), 400

        return jsonify({"inserted": len(rates)}), 201

    # ---------------------------------------------------------------
    # Trucks
    # ---------------------------------------------------------------
    @app.route("/truck", methods=["POST"])
    def create_truck():
        return jsonify({"id": "TRUCK123", "provider": 0}), 201

    @app.route("/truck/<truck_id>", methods=["PUT"])
    def update_truck(truck_id: str):
        return jsonify({"id": truck_id, "provider": 0}), 200

    @app.route("/truck/<truck_id>", methods=["GET"])
    def get_truck_info(truck_id: str):
        return jsonify({
            "id": truck_id,
            "tara": 0,
            "sessions": [],
            "placeholder": True
        }), 200

    # ---------------------------------------------------------------
    # Billing
    # ---------------------------------------------------------------
    @app.route("/bill/<int:provider_id>", methods=["GET"])
    def get_bill(provider_id: int):
        return jsonify({
            "id": str(provider_id),
            "name": "placeholder-provider",
            "from": "YYYYMMDDHHMMSS",
            "to": "YYYYMMDDHHMMSS",
            "truckCount": 0,
            "sessionCount": 0,
            "products": [],
            "total": 0,
            "placeholder": True
        }), 200

    return app


if __name__ == "__main__":
    app = create_app()
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)

